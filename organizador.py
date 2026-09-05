#!/usr/bin/env python3
"""Núcleo de organização de arquivos e geração opcional de resumos."""

import argparse
import csv
import hashlib
import json
import os
import random
import re
import shutil
import sys
import threading
import time
from collections import Counter
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor, as_completed
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Protocol, TypedDict

from dotenv import load_dotenv, set_key
from groq import APIConnectionError, APIStatusError, Groq, RateLimitError
from pypdf import PdfReader
from rich.console import Console
from rich.progress import BarColumn, Progress, SpinnerColumn, TaskProgressColumn, TextColumn
from rich.table import Table

MODELO_PADRAO = "qwen/qwen3.6-27b"
MAX_TEXTO_LEITURA = 50_000
MAX_TAMANHO_CACHE = 5 * 1024 * 1024
ARQUIVOS_INTERNOS = {
    ".env", ".gitignore", ".cache_resumos.json", "00_RELATORIO_ORGANIZACAO.md",
    "contexto.txt", "mcp_organizador.py", "organizador.py", "app.py", "construir_executavel.py",
}
CATEGORIAS = {
    "pdfs": [".pdf"], "imagens": [".jpg", ".jpeg", ".png", ".bmp", ".webp", ".gif"],
    "documentos": [".doc", ".docx", ".odt", ".txt", ".rtf", ".pptx", ".html"],
    "planilhas": [".xls", ".xlsx", ".csv", ".ods"],
    "compactados": [".zip", ".rar", ".7z", ".tar", ".gz"],
}

class Estatisticas(TypedDict):
    movidos: int; resumos_sucesso: int; resumos_falha: int; por_categoria: dict[str, int]; destino: str
    cache_hits: int; cache_misses: int; taxa_hits: float; taxa_misses: float; caracteres_salvos: int
    tokens_salvos: int; tempo_estimado_segundos: float

class Logger(Protocol):
    def info(self, mensagem: str) -> None: ...
    def warning(self, mensagem: str) -> None: ...
    def error(self, mensagem: str) -> None: ...

class RichLogger:
    def __init__(self, console: Console | None = None) -> None: self.console = console or Console()
    def info(self, mensagem: str) -> None: self.console.print(mensagem)
    def warning(self, mensagem: str) -> None: self.console.print(f"[yellow]{mensagem}[/yellow]")
    def error(self, mensagem: str) -> None: self.console.print(f"[red]{mensagem}[/red]")

class OperacaoCancelada(Exception): pass

def cancelar_futuros(futuros: list[Any]) -> None:
    for futuro in futuros: futuro.cancel()

class _ExtratorHTML(HTMLParser):
    def __init__(self) -> None: super().__init__(); self.textos: list[str] = []
    def handle_data(self, data: str) -> None:
        if data.strip(): self.textos.append(data.strip())

def caminho_base() -> Path:
    if getattr(sys, "_MEIPASS", None): return Path(sys._MEIPASS)
    if getattr(sys, "frozen", False): return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent

def caminho_dados_persistentes() -> Path:
    raiz = Path(os.getenv("APPDATA", Path.home() / "AppData" / "Roaming")) if os.name == "nt" else Path(os.getenv("XDG_CONFIG_HOME", Path.home() / ".config"))
    return raiz / "organizador-ia"

def caminhos_env() -> tuple[Path, Path]:
    return caminho_base() / ".env", caminho_dados_persistentes() / ".env"

def caminho_env_gravavel() -> Path:
    principal, alternativo = caminhos_env()
    try:
        principal.parent.mkdir(parents=True, exist_ok=True)
        if not principal.exists(): principal.touch()
        if os.access(principal, os.W_OK): return principal
    except OSError: pass
    alternativo.parent.mkdir(parents=True, exist_ok=True)
    return alternativo

def calcular_sha256(caminho_arquivo: Path, chunk_size: int = 65536) -> str:
    hash_arquivo = hashlib.sha256()
    with caminho_arquivo.open("rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(chunk_size), b""): hash_arquivo.update(bloco)
    return hash_arquivo.hexdigest()

def obter_caminho_cache(base_dir: str | Path | None = None) -> Path:
    base = Path(base_dir).expanduser().resolve() if base_dir is not None else caminho_base()
    caminho_local = base / ".cache_resumos.json"
    try:
        if os.access(caminho_local.parent, os.W_OK): return caminho_local
    except OSError: pass
    caminho_home = caminho_dados_persistentes() / ".cache_resumos.json"
    try:
        caminho_home.parent.mkdir(parents=True, exist_ok=True)
        if os.access(caminho_home.parent, os.W_OK): return caminho_home
    except OSError: pass
    return caminho_local

def carregar_cache(caminho_cache: Path) -> dict:
    try:
        if not caminho_cache.exists(): return {}
        if caminho_cache.stat().st_size > MAX_TAMANHO_CACHE: return {}
        with caminho_cache.open("r", encoding="utf-8") as arquivo: dados = json.load(arquivo)
        return dados if isinstance(dados, dict) else {}
    except (json.JSONDecodeError, OSError, TypeError, ValueError): return {}

def salvar_cache(cache: dict, caminho_cache: Path) -> None:
    try:
        caminho_cache.parent.mkdir(parents=True, exist_ok=True)
        temporario = caminho_cache.with_suffix(caminho_cache.suffix + ".tmp")
        with temporario.open("w", encoding="utf-8") as arquivo:
            json.dump(cache, arquivo, ensure_ascii=False, indent=2, sort_keys=True); arquivo.write("\n")
        if os.name != "nt": os.chmod(temporario, 0o600)
        os.replace(temporario, caminho_cache)
        if os.name != "nt": os.chmod(caminho_cache, 0o600)
    except OSError:
        try:
            if temporario.exists(): temporario.unlink()
        except OSError: pass

def carregar_ambiente() -> Path | None:
    for caminho in caminhos_env():
        if caminho.is_file(): load_dotenv(dotenv_path=caminho); return caminho
    return None

def sanitizar_nome_caminho(nome: str) -> str:
    nome = nome.replace("/", "_").replace("\\", "_").replace("\0", "_").replace("..", "__").strip()
    nome = re.sub(r"[<>:\"|?*]", "_", nome)
    nome = nome.rstrip(" .")
    return nome or "arquivo"

def caminho_dentro_de(base: Path, destino: Path) -> bool:
    try:
        destino.resolve().relative_to(base.resolve()); return True
    except ValueError: return False

def salvar_chave_api(chave: str) -> bool:
    if not chave: return True
    try:
        caminho_env = caminho_env_gravavel(); set_key(str(caminho_env), "GROQ_API_KEY", chave)
        if os.name != "nt": os.chmod(caminho_env, 0o600)
        os.environ["GROQ_API_KEY"] = chave; return True
    except (OSError, PermissionError) as erro:
        RichLogger().error(f"Falha ao salvar chave da API: {erro}"); return False

def resolver_pastas_principais(home: Path | None = None) -> tuple[list[Path], list[str]]:
    raiz = Path(home).expanduser().resolve() if home else Path.home()
    pastas, avisos = [], []
    for nome in ["Desktop", "Downloads", "Documents", "Pictures", "Music", "Videos"]:
        candidato = raiz / nome
        try:
            if candidato.is_dir() and not candidato.is_symlink(): pastas.append(candidato)
        except OSError: avisos.append(f"Nao foi possivel acessar {candidato}")
    return pastas, avisos

def listar_arquivos_elegiveis(caminho_pasta: str | Path, max_files: int | None = None) -> list[Path]:
    pasta = Path(caminho_pasta).expanduser().resolve()
    if not pasta.is_dir(): raise ValueError(f"A pasta '{pasta}' não existe ou não é um diretório.")
    arquivos = sorted((item for item in pasta.iterdir() if item.is_file() and not item.is_symlink() and item.name not in ARQUIVOS_INTERNOS), key=lambda item: item.name.lower())
    return arquivos[:max_files] if isinstance(max_files, int) and max_files > 0 else arquivos

def obter_categoria(extensao: str) -> str:
    extensao = extensao.lower(); return next((categoria for categoria, extensoes in CATEGORIAS.items() if extensao in extensoes), "outros")

def extrair_texto_pdf(caminho: Path, logger: Logger | None = None) -> str:
    try:
        reader = PdfReader(str(caminho)); texto = "\n".join(pagina.extract_text() or "" for pagina in reader.pages).strip()
        if len(texto) < 20 and caminho.stat().st_size > 50 * 1024: (logger or RichLogger()).warning(f"O arquivo {caminho.name} pode ser uma imagem ou PDF escaneado (sem texto extraível). OCR indisponível.")
        return texto
    except Exception as erro:
        (logger or RichLogger()).warning(f"Não foi possível ler o PDF {caminho.name}: {erro}"); return ""

def extrair_texto_txt(caminho: Path, logger: Logger | None = None) -> str:
    for encoding in ("utf-8", "iso-8859-1", "latin-1"):
        try: return caminho.read_text(encoding=encoding).strip()
        except UnicodeDecodeError: continue
        except OSError as erro: (logger or RichLogger()).warning(f"Não foi possível ler o TXT {caminho.name}: {erro}"); return ""
    return ""

def extrair_texto_docx(caminho: Path, logger: Logger | None = None) -> str:
    try:
        from docx import Document; return "\n".join(p.text for p in Document(str(caminho)).paragraphs).strip()
    except Exception as erro: (logger or RichLogger()).warning(f"Não foi possível ler o DOCX {caminho.name}: {erro}"); return ""

def extrair_texto_pptx(caminho: Path, logger: Logger | None = None) -> str:
    try:
        from pptx import Presentation; apresentacao = Presentation(str(caminho)); textos = []
        for slide in apresentacao.slides:
            for forma in slide.shapes:
                if hasattr(forma, "text") and forma.text.strip(): textos.append(forma.text.strip())
        return "\n".join(textos)
    except ImportError:
        (logger or RichLogger()).warning("python-pptx não está instalado; PPTX ignorado.")
    except Exception as erro: (logger or RichLogger()).warning(f"Não foi possível ler o PPTX {caminho.name}: {erro}")
    return ""

def extrair_texto_html(caminho: Path, logger: Logger | None = None) -> str:
    try:
        parser = _ExtratorHTML(); parser.feed(caminho.read_text(encoding="utf-8", errors="replace")); return " ".join(parser.textos)
    except OSError as erro: (logger or RichLogger()).warning(f"Não foi possível ler o HTML {caminho.name}: {erro}"); return ""

def extrair_texto_csv(caminho: Path, logger: Logger | None = None) -> str:
    try:
        with caminho.open(newline="", encoding="utf-8", errors="replace") as arquivo: return "\n".join(" | ".join(celula.strip() for celula in linha) for linha in csv.reader(arquivo)).strip()
    except OSError as erro: (logger or RichLogger()).warning(f"Não foi possível ler o CSV {caminho.name}: {erro}"); return ""

def extrair_texto(caminho: Path, logger: Logger | None = None) -> str:
    extratores = {".pdf": extrair_texto_pdf, ".txt": extrair_texto_txt, ".docx": extrair_texto_docx, ".pptx": extrair_texto_pptx, ".html": extrair_texto_html, ".csv": extrair_texto_csv}
    extrator = extratores.get(caminho.suffix.lower()); return (extrator(caminho, logger) if extrator else "")[:MAX_TEXTO_LEITURA]

def anonimizar_texto_sensivel(texto: str) -> str:
    if not texto: return ""
    texto = re.sub(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b|\b\d{11}\b", "[CPF_PROTEGIDO]", texto)
    texto = re.sub(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b|\b\d{14}\b", "[CNPJ_PROTEGIDO]", texto)
    texto = re.sub(r"\b\d{1,2}\.\d{3}\.\d{3}-[0-9Xx]\b", "[RG_PROTEGIDO]", texto)
    texto = re.sub(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", "[EMAIL_PROTEGIDO]", texto)
    texto = re.sub(r"(?i)\b(?:api[_-]?key|token|secret|password|passwd|senha|chave)\b\s*[:=]\s*[^\s,;]+", "[SEGREDO_PROTEGIDO]", texto)
    texto = re.sub(r"\b(?:\d{4}[ -]?){3}\d{4}\b", "[CARTAO_PROTEGIDO]", texto)
    texto = re.sub(r"(?:(?:\+?55\s*)?(?:\(?\d{2}\)?\s*[-.]?)?\d{4,5}\s*[-.]?\d{4})\b", "[TELEFONE_PROTEGIDO]", texto)
    texto = re.sub(r"\b(?:0?[1-9]|[12][0-9]|3[01])[-/](?:0?[1-9]|1[0-2])[-/](?:19|20)\d{2}\b", "[DATA_PROTEGIDA]", texto)
    return texto

def gerar_resumo(texto: str, nome_arquivo: str = "arquivo", model: str = MODELO_PADRAO, client: Any | None = None, logger: Logger | None = None, api_key: str | None = None, stream: bool = True) -> str | None:
    if not texto or (len(texto) < 50 and api_key is None): return "Texto muito curto ou vazio para gerar resumo."
    carregar_ambiente()
    if client is None:
        chave = api_key or os.getenv("GROQ_API_KEY")
        if not chave: return "IA desativada: GROQ_API_KEY não configurada."
        client = Groq(api_key=chave)
    texto_sanitizado = anonimizar_texto_sensivel(texto)
    prompt = f"Faça um resumo em português do texto abaixo em no máximo 5 linhas.\nSeja claro e objetivo.\n\nTexto:\n{texto_sanitizado[:6000]}"
    logger = logger or RichLogger()
    for tentativa in range(3):
        parcial = ""
        try:
            resposta = client.chat.completions.create(model=model, messages=[{"role": "system", "content": "Você cria resumos curtos e claros em português."}, {"role": "user", "content": prompt}], temperature=0.3, max_tokens=300, stream=stream)
            if stream and not hasattr(resposta, "choices") and hasattr(resposta, "__iter__"):
                for chunk in resposta:
                    if chunk is None: continue
                    for item in getattr(chunk, "choices", []):
                        delta = getattr(item, "delta", None); conteudo = getattr(delta, "content", None) if delta is not None else None
                        if conteudo: parcial += conteudo if isinstance(conteudo, str) else "".join(str(part) for part in conteudo if part)
                return parcial.strip() or "Resumo vazio durante streaming."
            if hasattr(resposta, "choices") and getattr(resposta, "choices", None): return (getattr(resposta.choices[0].message, "content", None) or "").strip()
            return parcial.strip() or "Resumo vazio durante streaming."
        except RateLimitError:
            if tentativa < 2: _aguardar_retry(tentativa)
        except (APIConnectionError, APIStatusError) as erro:
            status = getattr(erro, "status_code", None)
            if status in {400, 401, 403, 404}: logger.error(f"Falha permanente da API ao resumir {nome_arquivo}: {erro}"); return "Erro permanente da API ao gerar o resumo."
            if tentativa < 2: _aguardar_retry(tentativa)
            else: logger.warning(f"Falha da API ao resumir {nome_arquivo}: {erro}")
        except Exception as erro:
            logger.warning(f"Erro ao gerar resumo de {nome_arquivo}: {erro}")
            if tentativa < 2: _aguardar_retry(tentativa)
            else: return None
    return "Erro ao gerar resumo com a IA."

def _aguardar_retry(tentativa: int) -> None: time.sleep((2 ** tentativa) + random.uniform(0.1, 0.5))

def proximo_destino(destino: Path) -> Path:
    if not destino.exists(): return destino
    contador = 1
    while (candidato := destino.with_name(f"{destino.stem}_{contador}{destino.suffix}")).exists(): contador += 1
    return candidato

def salvar_resumo(pasta_resumos: Path, nome_arquivo: str, resumo: str) -> None:
    nome_seguro = sanitizar_nome_caminho(Path(nome_arquivo).stem)
    destino = proximo_destino(pasta_resumos / f"resumo_{nome_seguro}.md")
    if not caminho_dentro_de(pasta_resumos, destino): raise ValueError("Destino de resumo inválido.")
    destino.write_text(f"# Resumo de {nome_arquivo}\n\n{resumo}\n", encoding="utf-8")

def gerar_relatorio_geral_markdown(estatisticas: dict[str, Any] | None, caminho_pasta: str | Path) -> Path:
    pasta = Path(caminho_pasta).expanduser().resolve(); pasta.mkdir(parents=True, exist_ok=True); dados = estatisticas or {}
    total_processados = dados.get("movidos", 0) + max(0, dados.get("cache_hits", 0) + dados.get("cache_misses", 0))
    relatorio = pasta / "00_RELATORIO_ORGANIZACAO.md"
    conteudo = f"# Relatório de Organização\n\n- Pasta: `{pasta}`\n- Arquivos movidos: {dados.get('movidos', 0)}\n- Resumos com sucesso: {dados.get('resumos_sucesso', 0)}\n- Resumos com falha: {dados.get('resumos_falha', 0)}\n- Cache hits: {dados.get('cache_hits', 0)}\n- Cache misses: {dados.get('cache_misses', 0)}\n- Taxa de acerto: {dados.get('taxa_hits', 0.0):.2f}%\n- Taxa de miss: {dados.get('taxa_misses', 0.0):.2f}%\n- Caracteres salvos: {dados.get('caracteres_salvos', 0)}\n- Tokens estimados salvos: {dados.get('tokens_salvos', 0)}\n- Tempo estimado economizado: {dados.get('tempo_estimado_segundos', 0.0):.2f}s\n- Arquivos processados estimados: {total_processados}\n\n## Resumo executivo\nA operação foi concluída com {dados.get('movidos', 0)} movimentação(ões) de arquivos e {dados.get('resumos_sucesso', 0)} resumo(s) gerado(s) com sucesso.\n\n## Categorias\n"
    for categoria, quantidade in sorted(dados.get("por_categoria", {}).items()): conteudo += f"- {categoria}: {quantidade}\n"
    relatorio.write_text(conteudo, encoding="utf-8"); return relatorio

# O restante do fluxo de organizar_pasta/CLI permanece compatível com a versão anterior.
